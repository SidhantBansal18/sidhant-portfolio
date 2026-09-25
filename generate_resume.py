from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_resume():
    doc = Document()
    
    # Set global style
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    # Header
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run('SIDHANT BANSAL')
    run.bold = True
    run.font.size = Pt(16)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.add_run('Seattle, WA | +1 (425) 382-9156 | sidhantbansal18@outlook.com | linkedin.com/in/sidhantbansal18')

    # Technical Skills
    doc.add_heading('TECHNICAL SKILLS', level=1)
    skills = doc.add_paragraph()
    skills.add_run('Languages: ').bold = True
    skills.add_run('Python, Java, C++, C#, JavaScript, TypeScript, SQL\n')
    skills.add_run('Frameworks & Cloud: ').bold = True
    skills.add_run('React, Node.js, Express, Spring Boot, ASP.NET, AWS, Azure, Docker, Laravel, Angular\n')
    skills.add_run('Databases & Tools: ').bold = True
    skills.add_run('PostgreSQL, MySQL, MongoDB, SQL Server, Git, JIRA, ArcGIS, Unity, GitHub Copilot')

    # Experience
    doc.add_heading('PROFESSIONAL EXPERIENCE', level=1)
    
    # Washington State Dept of Ecology
    exp1 = doc.add_paragraph()
    run1 = exp1.add_run('Washington State Department of Ecology')
    run1.bold = True
    exp1.add_run(' | Seattle, WA')
    
    role1 = doc.add_paragraph()
    role1.add_run('Software Engineer').italic = True
    role1.add_run(' | Aug 2023 – Present')

    # Maritime Risk Model
    doc.add_paragraph('Maritime Risk Model & Data Infrastructure', style='List Bullet')
    doc.add_paragraph('Built a high-performance ETL pipeline in Python to process 10M+ AIS messages, simulating vessel traffic and movement patterns to inform state safety policy.', style='List Bullet')
    doc.add_paragraph('Optimized data ingestion by migrating to Polars, reducing extraction time for million-row datasets by 80%.', style='List Bullet')
    doc.add_paragraph('Engineered a momentum and drift simulation module using Numba and Polars; implemented a 3D matrix for weather data to enable instant lookups, reducing total simulation time from 2 hours to 5 minutes.', style='List Bullet')
    doc.add_paragraph('Improved database write performance by implementing SQLAlchemy chunking and converting geometry columns from WKT to WKB format.', style='List Bullet')
    doc.add_paragraph('Developed algorithms to detect vessel "laden" status and tug escort events, identifying 2,000+ maritime incidents for risk assessment.', style='List Bullet')

    # SPIIS
    doc.add_paragraph('SPIIS (Environmental Incident Management Platform)', style='List Bullet')
    doc.add_paragraph('Developed and maintained the SPIIS platform using ASP.NET and C#, reducing feature implementation time by 30% through AI-assisted development workflows.', style='List Bullet')
    doc.add_paragraph('Optimized the database layer by implementing stored procedures and normalizing schemas (one-to-many relationships), reducing feature load times by 50%.', style='List Bullet')
    doc.add_paragraph('Collaborated with state stakeholders to translate environmental safety requirements into technical specifications and production features.', style='List Bullet')

    # Platform Science
    exp2 = doc.add_paragraph()
    run2 = exp2.add_run('Platform Science')
    run2.bold = True
    exp2.add_run(' | San Diego, CA')
    
    role2 = doc.add_paragraph()
    role2.add_run('Software Developer Intern').italic = True
    role2.add_run(' | Jun 2022 – Sep 2022')

    doc.add_paragraph('Built RESTful APIs using Node.js, Express, and Laravel (PHP) with a focus on service versioning and backward compatibility.', style='List Bullet')
    doc.add_paragraph('Reduced manual testing overhead by introducing automated unit tests via Docker-based CI, increasing overall development velocity by 15%.', style='List Bullet')
    doc.add_paragraph('Refactored TypeScript and PHP codebases using modular design patterns to improve component reusability.', style='List Bullet')

    # Accenture
    exp3 = doc.add_paragraph()
    run3 = exp3.add_run('Accenture')
    run3.bold = True
    exp3.add_run(' | Bangalore, India')
    
    role3 = doc.add_paragraph()
    role3.add_run('Application Developer').italic = True
    role3.add_run(' | Sep 2019 – Aug 2021')

    doc.add_paragraph('Developed full-stack web applications using the MEAN stack (MongoDB, Express, Angular, Node.js), supporting a 30% increase in concurrent user capacity.', style='List Bullet')
    doc.add_paragraph('Reduced API response times by 40% by implementing optimized database indexing and query tuning.', style='List Bullet')
    doc.add_paragraph('Partnered with PMs in an Agile environment to ship UI/UX optimizations that increased user engagement by 25%.', style='List Bullet')

    # Projects
    doc.add_heading('KEY PROJECTS', level=1)
    
    # VR Platform
    p1 = doc.add_paragraph()
    runp1 = p1.add_run('VR Therapy Platform')
    runp1.bold = True
    p1.add_run(' | Graduate Research (EYE Research Group) | Sep 2022 – Jun 2023')
    doc.add_paragraph('Consolidated four disparate VR applications into a unified Unity/C# platform, reducing code redundancy by 35% through modular refactoring.', style='List Bullet')
    doc.add_paragraph('Optimized rendering pipelines to improve frame rate and performance by 20%.', style='List Bullet')

    # Dependify
    p2 = doc.add_paragraph()
    runp2 = p2.add_run('Dependify')
    runp2.bold = True
    p2.add_run(' | Distributed Systems Project | Apr 2022 – Jun 2022')
    doc.add_paragraph('Architected a microservices system using Java, Spring Boot, and PostgreSQL on AWS EC2 to manage software dependencies.', style='List Bullet')
    doc.add_paragraph('Designed the system to handle 10,000+ daily requests with a 40% improvement in response latency compared to monolithic baselines.', style='List Bullet')

    # Education
    doc.add_heading('EDUCATION', level=1)
    
    edu1 = doc.add_paragraph()
    runedu1 = edu1.add_run('University of Washington')
    runedu1.bold = True
    edu1.add_run(' | Seattle, WA')
    
    deg1 = doc.add_paragraph()
    deg1.add_run('Master of Science in Computer Science').italic = True
    deg1.add_run(' | Jun 2023')
    doc.add_paragraph('Relevant Coursework: Parallel Computing, Distributed Computing, Algorithms, Computer Graphics', style='List Bullet')

    edu2 = doc.add_paragraph()
    runedu2 = edu2.add_run('Guru Gobind Singh Indraprastha University')
    runedu2.bold = True
    edu2.add_run(' | New Delhi, India')
    
    deg2 = doc.add_paragraph()
    deg2.add_run('Bachelor of Technology in Information Technology').italic = True
    deg2.add_run(' | Jun 2019')
    doc.add_paragraph('Relevant Coursework: Data Structures and Algorithms, Operating Systems, Cloud Computing', style='List Bullet')

    doc.save('Sidhant_Bansal_Resume_Optimized.docx')

create_resume()
