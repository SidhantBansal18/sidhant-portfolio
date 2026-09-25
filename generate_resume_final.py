from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_resume():
    doc = Document()
    
    # Set margins
    sections = doc.sections[0]
    sections.top_margin = Inches(0.5)
    sections.bottom_margin = Inches(0.5)
    sections.left_margin = Inches(0.5)
    sections.right_margin = Inches(0.5)

    # Global style
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)

    # Header
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run('SIDHANT BANSAL')
    run.bold = True
    run.font.size = Pt(14)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.add_run('Seattle, WA | +1 (425) 382-9156 | sidhantbansal18@outlook.com | linkedin.com/in/sidhantbansal18')
    contact.paragraph_format.space_after = Pt(6)

    def add_section_heading(text):
        h = doc.add_paragraph()
        run = h.add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(3)

    # Technical Skills
    add_section_heading('TECHNICAL SKILLS')
    skills = doc.add_paragraph()
    skills.add_run('Languages: ').bold = True
    skills.add_run('Python, Java, C++, C#, JavaScript, TypeScript, SQL, HTML/CSS\n')
    skills.add_run('Frameworks & Cloud: ').bold = True
    skills.add_run('React, Node.js, Express, Spring Boot, ASP.NET Core, AWS (EC2, S3), Azure DevOps, Docker, Laravel, Angular\n')
    skills.add_run('Databases & Tools: ').bold = True
    skills.add_run('PostgreSQL, MySQL, MongoDB, SQL Server, Redis, Git, JIRA, ArcGIS, Unity, GitHub Copilot, SQLAlchemy')
    skills.paragraph_format.space_after = Pt(6)

    # Experience
    add_section_heading('PROFESSIONAL EXPERIENCE')
    
    # Washington State Dept of Ecology
    exp1 = doc.add_paragraph()
    run1 = exp1.add_run('Washington State Department of Ecology')
    run1.bold = True
    exp1.add_run(' | Seattle, WA')
    
    role1 = doc.add_paragraph()
    role1.add_run('Software Engineer').italic = True
    role1.add_run(' | Aug 2023 – Present')

    doc.add_paragraph('Architected a high-performance ETL pipeline in Python to process 10M+ AIS messages; migrated data ingestion from Pandas to Polars to achieve an 80% reduction in extraction time for million-row datasets.', style='List Bullet')
    doc.add_paragraph('Engineered a momentum and drift simulation module using Numba (JIT compilation) and Polars; implemented a 3D matrix for weather data to enable O(1) lookups, reducing total simulation runtime from 2 hours to 5 minutes.', style='List Bullet')
    doc.add_paragraph('Optimized database write throughput by implementing SQLAlchemy chunking and migrating geometry columns from WKT (Well-Known Text) to WKB (Well-Known Binary) format for faster spatial indexing.', style='List Bullet')
    doc.add_paragraph('Developed a suite of algorithms to detect vessel "laden" status and tug escort events based on movement patterns and draft analysis, identifying 2,000+ maritime incidents to inform state safety policy.', style='List Bullet')
    doc.add_paragraph('Maintained a mission-critical ASP.NET/C# platform for incident management; reduced feature implementation time by 30% through AI-assisted development and optimized database performance by 50% via stored procedure tuning and schema normalization.', style='List Bullet')
    doc.add_paragraph('Facilitated monthly brainstorming sessions with state program leads to identify operational bottlenecks and translate safety requirements into technical tickets and production features.', style='List Bullet')

    # Platform Science
    exp2 = doc.add_paragraph()
    run2 = exp2.add_run('Platform Science')
    run2.bold = True
    exp2.add_run(' | San Diego, CA')
    
    role2 = doc.add_paragraph()
    role2.add_run('Software Developer Intern').italic = True
    role2.add_run(' | Jun 2022 – Sep 2022')

    doc.add_paragraph('Engineered RESTful APIs using Node.js, Express, and Laravel (PHP), ensuring strict backward compatibility and service versioning for legacy client support.', style='List Bullet')
    doc.add_paragraph('Optimized the development lifecycle by introducing automated unit testing and Docker-based CI/CD pipelines, increasing development velocity by 15%.', style='List Bullet')
    doc.add_paragraph('Refactored TypeScript and PHP codebases using modular design patterns (Strategy and Factory patterns) to reduce code duplication and improve component reusability.', style='List Bullet')

    # Accenture
    exp3 = doc.add_paragraph()
    run3 = exp3.add_run('Accenture')
    run3.bold = True
    exp3.add_run(' | Bangalore, India')
    
    role3 = doc.add_paragraph()
    role3.add_run('Application Developer').italic = True
    role3.add_run(' | Sep 2019 – Aug 2021')

    doc.add_paragraph('Developed and scaled full-stack web applications using the MEAN stack (MongoDB, Express, Angular, Node.js), implementing caching strategies that supported a 30% increase in concurrent user capacity.', style='List Bullet')
    doc.add_paragraph('Reduced API response times by 40% through the implementation of optimized MongoDB indexing and query profiling to eliminate bottlenecks in data retrieval.', style='List Bullet')
    doc.add_paragraph('Partnered with Product Managers in an Agile environment to deploy UI/UX optimizations and feature enhancements that increased user engagement by 25%.', style='List Bullet')

    # Projects
    add_section_heading('KEY PROJECTS')
    
    # VR Platform
    p1 = doc.add_paragraph()
    runp1 = p1.add_run('VR Therapy Platform')
    runp1.bold = True
    p1.add_run(' | Graduate Research (EYE Research Group) | Sep 2022 – Jun 2023')
    doc.add_paragraph('Unified four separate VR therapy applications into a single, modular Unity/C# platform, reducing code redundancy by 35% and improving rendering performance by 20% through optimized draw calls.', style='List Bullet')

    # Dependify
    p2 = doc.add_paragraph()
    runp2 = p2.add_run('Dependify')
    runp2.bold = True
    p2.add_run(' | Distributed Systems Project | Apr 2022 – Jun 2022')
    doc.add_paragraph('Architected a distributed microservices system using Java, Spring Boot, and PostgreSQL hosted on AWS EC2 for centralized dependency management.', style='List Bullet')
    doc.add_paragraph('Designed the system to handle 10,000+ daily requests, achieving a 40% reduction in response latency compared to traditional monolithic architectures.', style='List Bullet')

    # Education
    add_section_heading('EDUCATION')
    
    edu1 = doc.add_paragraph()
    runedu1 = edu1.add_run('University of Washington')
    runedu1.bold = True
    edu1.add_run(' | Seattle, WA | Master of Science in Computer Science | Jun 2023')
    doc.add_paragraph('Coursework: Parallel Computing, Distributed Computing, Algorithms, Computer Graphics, Software Management', style='List Bullet')

    edu2 = doc.add_paragraph()
    runedu2 = edu2.add_run('Guru Gobind Singh Indraprastha University')
    runedu2.bold = True
    edu2.add_run(' | New Delhi, India | B.Tech in Information Technology | Jun 2019')
    doc.add_paragraph('Coursework: Data Structures and Algorithms, Operating Systems, Cloud Computing, Computer Architecture', style='List Bullet')

    doc.save('Sidhant_Bansal_Resume_Final_Expanded.docx')

create_resume()
