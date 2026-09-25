from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

def create_resume():
    doc = Document()
    
    # Tighten margins to fit everything on 1 page
    sections = doc.sections[0]
    sections.top_margin = Inches(0.5)
    sections.bottom_margin = Inches(0.5)
    sections.left_margin = Inches(0.5)
    sections.right_margin = Inches(0.5)

    # Global style
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)

    # Header
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run('SIDHANT BANSAL')
    run.bold = True
    run.font.size = Pt(14)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.add_run('Seattle, WA | +1 (425) 382-9156 | sidhantbansal18@outlook.com | linkedin.com/in/sidhantbansal18')
    contact.paragraph_format.space_after = Pt(6)

    def add_section_heading(text):
        h = doc.add_paragraph()
        run = h.add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        # Add a bottom border manually by drawing a line if possible, 
        # but for simple ATS, just bold + small spacing works best.
        h.paragraph_format.space_before = Pt(6)
        h.paragraph_format.space_after = Pt(2)

    # Technical Skills
    add_section_heading('TECHNICAL SKILLS')
    skills = doc.add_paragraph()
    skills.add_run('Languages: ').bold = True
    skills.add_run('Python, Java, C++, C#, JavaScript, TypeScript, SQL\n')
    skills.add_run('Frameworks & Cloud: ').bold = True
    skills.add_run('React, Node.js, Express, Spring Boot, ASP.NET, AWS, Azure, Docker, Laravel, Angular\n')
    skills.add_run('Databases & Tools: ').bold = True
    skills.add_run('PostgreSQL, MySQL, MongoDB, SQL Server, Git, JIRA, ArcGIS, Unity, GitHub Copilot')
    skills.paragraph_format.space_after = Pt(6)

    # Experience
    add_section_heading('PROFESSIONAL EXPERIENCE')
    
    # Washington State Dept of Ecology
    exp1 = doc.add_paragraph()
    run1 = exp1.add_run('Washington State Department of Ecology')
    run1.bold = True
    exp1.add_run(' | Seattle, WA')
    
    role1 = doc.add_paragraph()
    role1.add_run('Software Engineer').italic = True
    role1.add_run(' | Aug 2023 – Present')

    # Combine subtitles into bullets to save vertical space
    doc.add_paragraph('Maritime Risk Model: Built high-performance ETL pipeline for 10M+ AIS messages; optimized ingestion via Polars (80% faster) and developed Numba/Polars simulation module with 3D weather matrix, reducing runtime from 2 hrs to 5 mins.', style='List Bullet')
    doc.add_paragraph('Data Infra: Improved DB write performance via SQLAlchemy chunking and WKB geometry conversion; developed algorithms to identify 2,000+ maritime incidents via laden/escort status.', style='List Bullet')
    doc.add_paragraph('SPIIS Platform: Maintained ASP.NET/C# platform, reducing implementation time by 30% via AI workflows and reducing feature load times by 50% through stored procedures and schema normalization.', style='List Bullet')
    doc.add_paragraph('Stakeholder Mgmt: Collaborated with state agencies to translate safety requirements into technical specifications and production features.', style='List Bullet')

    # Platform Science
    exp2 = doc.add_paragraph()
    run2 = exp2.add_run('Platform Science')
    run2.bold = True
    exp2.add_run(' | San Diego, CA')
    
    role2 = doc.add_paragraph()
    role2.add_run('Software Developer Intern').italic = True
    role2.add_run(' | Jun 2022 – Sep 2022')

    doc.add_paragraph('Built RESTful APIs using Node.js, Express, and Laravel (PHP) supporting service versioning and backward compatibility.', style='List Bullet')
    doc.add_paragraph('Introduced automated unit tests via Docker-based CI, increasing development velocity by 15%.', style='List Bullet')
    doc.add_paragraph('Refactored TypeScript/PHP codebases using modular design patterns to improve reusability.', style='List Bullet')

    # Accenture
    exp3 = doc.add_paragraph()
    run3 = exp3.add_run('Accenture')
    run3.bold = True
    exp3.add_run(' | Bangalore, India')
    
    role3 = doc.add_paragraph()
    role3.add_run('Application Developer').italic = True
    role3.add_run(' | Sep 2019 – Aug 2021')

    doc.add_paragraph('Developed full-stack MEAN stack apps, supporting a 30% increase in concurrent user capacity.', style='List Bullet')
    doc.add_paragraph('Reduced API response times by 40% via optimized database indexing and query tuning.', style='List Bullet')
    doc.add_paragraph('Partnered with PMs in Agile sprints to ship UI/UX optimizations that increased engagement by 25%.', style='List Bullet')

    # Projects
    add_section_heading('KEY PROJECTS')
    
    # VR Platform
    p1 = doc.add_paragraph()
    runp1 = p1.add_run('VR Therapy Platform')
    runp1.bold = True
    p1.add_run(' | Graduate Research | Sep 2022 – Jun 2023')
    doc.add_paragraph('Unified four VR apps into a Unity/C# platform, reducing redundancy by 35% and improving performance by 20%.', style='List Bullet')

    # Dependify
    p2 = doc.add_paragraph()
    runp2 = p2.add_run('Dependify')
    runp2.bold = True
    p2.add_run(' | Distributed Systems Project | Apr 2022 – Jun 2022')
    doc.add_paragraph('Architected Java/Spring Boot microservices on AWS EC2 handling 10k+ daily requests with 40% lower latency.', style='List Bullet')

    # Education
    add_section_heading('EDUCATION')
    
    edu1 = doc.add_paragraph()
    runedu1 = edu1.add_run('University of Washington')
    runedu1.bold = True
    edu1.add_run(' | MS in Computer Science | Jun 2023')
    
    edu2 = doc.add_paragraph()
    runedu2 = edu2.add_run('Guru Gobind Singh Indraprastha University')
    runedu2.bold = True
    edu2.add_run(' | B.Tech in Information Technology | Jun 2019')

    doc.save('Sidhant_Bansal_Resume_1Page.docx')

create_resume()
